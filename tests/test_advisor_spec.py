"""The advisor against Bamshad's written specification.

Bamshad supplied three things: a three-stage flow with a named deliverable at
the end of each stage, a suggested organization for the finished proposal, and
a nineteen-point framework of lenses for interrogating a research idea. This
file holds the prompt and the schema to all three.

These are structural tests, not behavioural ones. They cannot prove the model
follows an instruction; they prove the instruction is still present. That is
worth having on its own, because the prompt is one 70KB string that several
people edit, and a rule deleted in a rewrite is invisible until a professor
notices the bot stopped doing something. Two of the rules covered here were in
fact dropped by an earlier rewrite and only caught by tests like these.
"""
import pytest

import web_app

STABLE, _ = web_app._advisor_system_prompt({"name": "Jane", "proposal": {}})
TOOLS = {t["function"]["name"]: t["function"] for t in web_app._ADVISOR_TOOLS}
SAVE_PROPOSAL = TOOLS["save_proposal"]["parameters"]["properties"]


# ── The finished proposal carries Bamshad's suggested organization ──────────

#   Title, Abstract, Keywords / Introduction and Overview (background, hypotheses
#   and objectives, assumptions and delimitations, importance and benefits) /
#   Related Work / Research Design and Methodology / Plan of Work and Outcomes /
#   Conclusions and Future Work / References
REQUIRED_SECTIONS = [
    "abstract", "keywords", "background", "objectives", "hypotheses",
    "assumptions_delimitations", "related_work", "methodology",
    "expected_outcomes", "plan_of_work", "conclusions_future_work",
]


@pytest.mark.parametrize("section", REQUIRED_SECTIONS)
def test_the_proposal_has_every_section_the_organization_calls_for(section):
    assert section in web_app._PROPOSAL_FIELDS


@pytest.mark.parametrize("section", REQUIRED_SECTIONS)
def test_the_advisor_can_actually_save_each_one(section):
    """A field in the schema that save_proposal cannot accept is a section the
    advisor can never fill."""
    assert section in SAVE_PROPOSAL


@pytest.mark.parametrize("section", REQUIRED_SECTIONS)
def test_each_section_reaches_the_downloaded_document(section, tmp_path):
    """The .docx is the deliverable. A section saved but never rendered would
    be invisible in the only artifact that leaves the app."""
    from docx import Document
    marker = f"UNIQUE-MARKER-{section}"
    docx = web_app._build_proposal_docx("Jane Doe", {section: marker})
    path = tmp_path / "p.docx"
    path.write_bytes(docx)
    text = "\n".join(p.text for p in Document(str(path)).paragraphs)
    assert marker in text, f"{section} never reaches the .docx"


@pytest.mark.parametrize("template", ["advisor.html", "projects.html"])
def test_both_panels_can_render_every_section(template):
    """Each template keeps its own label map. A section saved by the advisor
    but absent from the map is written to the database and never shown, which
    looks exactly like the advisor having ignored the request."""
    from pathlib import Path
    html = (Path(web_app.__file__).parent / "templates" / template).read_text()
    missing = [f for f in web_app._PROPOSAL_FIELDS if f not in html]
    assert missing == [], f"{template} cannot render {missing}"


def test_importance_and_benefits_is_required_somewhere():
    """Folded into background rather than given its own section, so the
    requirement has to be visible in background's own instructions."""
    assert "IMPORTANCE AND BENEFITS" in SAVE_PROPOSAL["background"]["description"]


def test_references_are_generated_rather_than_written():
    """Bamshad lists References; we build them from the works the novelty
    search actually returned instead of asking the model to recall citations."""
    from docx import Document
    docx = web_app._build_proposal_docx(
        "Jane Doe", {"abstract": "A."},
        references=[{"title": "Contesting risk scores", "authors": ["Lee, A."], "year": 2024}])
    import io
    text = "\n".join(p.text for p in Document(io.BytesIO(docx)).paragraphs)
    assert "References" in text
    assert "Contesting risk scores" in text


# ── Stage deliverables ─────────────────────────────────────────────────────

def test_the_problem_statement_is_three_to_four_paragraphs():
    """Bamshad's stage 1 output is "a 3-4 paragraph clear and concise statement
    of a research problem including properly framed context and research
    objectives". Ours asked for one paragraph until this was checked."""
    described = SAVE_PROPOSAL["problem_statement"]["description"]
    assert "THREE TO FOUR PARAGRAPHS" in described
    assert "a single paragraph is not enough" in described


@pytest.mark.parametrize("element", [
    "background that motivates",       # motivation and real-world impact
    "field of study",                  # context
    "scope",                           # scope
    "how the problem has been addressed before",
    "research objectives",
])
def test_the_problem_statement_covers_each_element_of_the_rubric(element):
    assert element in SAVE_PROPOSAL["problem_statement"]["description"]


@pytest.mark.parametrize("question", [
    "clear and unambiguous",
    "importance, its real-world impact",
    "Is the context clear",
    "what gap remains",
    "specific novel approach",
])
def test_the_five_assessment_questions_are_in_the_prompt(question):
    """The rubric the statement is judged on, stated where the advisor drafts
    it rather than left implicit."""
    assert question in STABLE


def test_stage_four_opens_with_the_two_page_outline():
    """Bamshad's stage 2 deliverable is a two-page outline of seven elements,
    produced before sections are developed. We had no such checkpoint: the
    advisor went straight from the problem statement into section-by-section
    drafting, so the proposal had no whole-shape draft at any point."""
    assert "FIRST, THE OUTLINE" in STABLE
    assert "roughly two pages" in STABLE


@pytest.mark.parametrize("element", [
    "general problem and the motivation",
    "specific research problem in its particular context",
    "What others have done",
    "knowledge or research gap",
    "key research objectives",
    "research questions that would guide",
    "expected results",
])
def test_the_outline_covers_all_seven_elements(element):
    assert element in STABLE


def test_the_outline_is_not_saved_as_proposal_sections():
    """It is a sketch meant to be thrown away. Saving it would fill the panel
    with text that never went through the section-by-section work."""
    assert "DO NOT save proposal sections from the outline" in STABLE


def test_methods_are_mapped_to_each_research_question():
    """Bamshad's methodology has three parts, and the middle one — "specific
    methods related to each research question" — is the one proposals skip. A
    question with no method is unanswerable as written; a method answering no
    question is scope nobody needs."""
    assert "THE METHOD FOR EACH RESEARCH QUESTION OR HYPOTHESIS, mapped one to one" in STABLE
    assert "A method that answers no question is scope" in STABLE


def test_the_data_analysis_approach_is_required_separately():
    assert "the data analysis approach" in STABLE


def test_hypotheses_stay_optional():
    """"Research questions and/or hypotheses" — and this roster runs from
    neuroscience to screenwriting. A hypothesis invented to fill a heading is
    worse than an empty section."""
    described = SAVE_PROPOSAL["hypotheses"]["description"]
    assert "OPTIONAL" in described
    assert "empty section is better" in described
    assert "humanistic" in STABLE


def test_assumptions_and_delimitations_are_kept_apart():
    """They are different things. Assumptions are what must be true;
    delimitations are the boundaries drawn on purpose."""
    described = SAVE_PROPOSAL["assumptions_delimitations"]["description"]
    assert "ASSUMPTIONS" in described and "DELIMITATIONS" in described
    assert "most work" in described


def test_the_plan_of_work_carries_a_schedule_and_deliverables():
    """"Expected outcomes" alone never gave us either."""
    described = SAVE_PROPOSAL["plan_of_work"]["description"]
    assert "schedule" in described and "deliverable" in described


# ── The nineteen lenses ────────────────────────────────────────────────────

# One distinctive phrase per lens. Chosen to be specific enough that a lens
# deleted in a rewrite fails here rather than passing on an incidental word.
LENSES = {
    1:  "CORE PHENOMENON",
    2:  "RESEARCH CLAIM",
    3:  "deeper consequence hiding inside",
    4:  "SCOPE AND BOUNDARIES",
    5:  "MECHANISM",
    6:  "COMPETING EXPLANATIONS",
    7:  "which does the MOST WORK",
    8:  "GENUINE UNCERTAINTY",
    9:  "CAUSAL STRUCTURE",
    10: "OPERATIONALIZATION AND MEASUREMENT",
    11: "UNIT AND LEVEL OF ANALYSIS",
    12: "DISCRIMINATING EVIDENCE",
    13: "FALSIFIABILITY",
    14: "empirically indistinguishable",
    15: "CONTRADICTIONS AND ANOMALIES",
    16: "CONTRIBUTION",
    17: "SKEPTICAL REVIEWER",
    18: "EXPERT-ONLY KNOWLEDGE",
    19: "THE MISSING DIMENSION",
}


@pytest.mark.parametrize("number,phrase", sorted(LENSES.items()))
def test_every_lens_is_present(number, phrase):
    assert phrase in STABLE, f"lens {number} is missing from the prompt"


def test_the_lenses_are_silent():
    """The failure mode they introduce. A researcher should meet a colleague
    who asks unusually good questions, not a rubric applied to them."""
    assert "never spoken aloud" in STABLE
    assert "Never name a lens" in STABLE
    assert "LENS LEAK" in STABLE


def test_the_lenses_still_produce_one_question():
    """The lens list is the strongest pull toward interrogation in the whole
    prompt, so it has to restate the limit itself."""
    assert "never ask about more than one in a turn" in STABLE


def test_the_lenses_do_not_licence_supplying_answers():
    """Competing explanations and discriminating evidence both tempt the model
    to hand over candidate content, which is what the interview rule forbids
    and what breaks the collaborator matching."""
    assert "THE STAGE 1 CONSTRAINT BINDS HERE TOO" in STABLE
    assert 'not "could it be selection or drift?" but "what else could produce that pattern?"' in STABLE


def test_the_feasibility_lens_does_not_reopen_stage_one_methods():
    """Lens 14 asks what data could realistically answer this. Stage 1 forbids
    exactly that question, because asking what data they have while the problem
    is vague reshapes the problem to fit the data. The stage rule wins, and the
    prompt has to say which."""
    assert "SOME LENSES ARE STAGE-GATED, and the stage rules win" in STABLE
    assert '"what would count as evidence" is a Stage 1 question' in STABLE


def test_competing_explanations_are_checked_against_the_literature():
    """Bamshad's own note on that lens. The search is the evidence, not the
    advisor's guess."""
    assert "CHECKED AGAINST THE LITERATURE in Stage 2" in STABLE


def test_expert_only_questions_are_preferred_not_merely_allowed():
    """The lens that most changes what a turn is worth. The prompt already
    banned asking what could be searched; it never stated the positive form."""
    assert "PREFER THESE ABOVE ALL OTHERS" in STABLE
    assert "If you could find out yourself, find out yourself" in STABLE
    assert "unpublished observation" in STABLE or "unpublished observations" in STABLE


def test_the_send_check_enforces_the_expert_only_preference():
    """A rule stated once in a 70KB prompt is a rule that gets skipped. The
    send check is what runs on every message."""
    assert "could only {name} answer".replace("{name}", "Jane") in STABLE
